from __future__ import annotations

import base64
import logging
import re
import subprocess
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from jinja2 import TemplateSyntaxError, UndefinedError

logger = logging.getLogger(__name__)

_JINJA_PRINT_RE = re.compile(r"\{\{.*?\}\}", flags=re.DOTALL)
_XML_TAG_RE = re.compile(r"<[^>]+>")
_SAFE_JINJA_KEY_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_PLACEHOLDER_ALIASES = {
    "дата рождения": "birth_date",
    "пол": "gender",
}


def _run_libreoffice(command, timeout, error_message):
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
    except FileNotFoundError as exc:
        raise RuntimeError("LibreOffice is not installed or not found in PATH") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(error_message) from exc
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice command failed: {result.stderr.strip()}")
    return result


def _decode_signature(signature_base64: str) -> bytes:
    if "," in signature_base64:
        signature_base64 = signature_base64.split(",", 1)[1]
    return base64.b64decode(signature_base64)


def _normalize_template_expression(raw_expression):
    plain_expression = _XML_TAG_RE.sub("", raw_expression)
    if not plain_expression.startswith("{{") or not plain_expression.endswith("}}"):
        return raw_expression, None
    placeholder = re.sub(r"\s+", " ", plain_expression[2:-2]).strip()
    alias_key = placeholder.lower()
    mapped = _PLACEHOLDER_ALIASES.get(alias_key)
    if mapped and mapped != placeholder:
        return f"{{{{ {mapped} }}}}", placeholder
    if _SAFE_JINJA_KEY_RE.fullmatch(placeholder):
        return raw_expression, None
    return raw_expression, placeholder


def _prepare_template_for_render(template_path, output_dir, output_basename):
    source = _resolve_renderable_template(template_path=template_path, output_dir=output_dir)
    normalized_path = Path(output_dir) / f"{output_basename}_template.docx"
    rewritten_placeholders = []
    suspicious_placeholders = []
    with ZipFile(source, "r") as src, ZipFile(normalized_path, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                xml = data.decode("utf-8")
                def _replace(match):
                    rewritten, suspicious = _normalize_template_expression(match.group(0))
                    if rewritten != match.group(0):
                        rewritten_placeholders.append(suspicious or "")
                    elif suspicious:
                        suspicious_placeholders.append(suspicious)
                    return rewritten
                xml = _JINJA_PRINT_RE.sub(_replace, xml)
                data = xml.encode("utf-8")
            dst.writestr(info, data)
    return normalized_path, sorted(set(suspicious_placeholders))


def _resolve_renderable_template(template_path, output_dir):
    source = Path(template_path)
    suffix = source.suffix.lower()
    if suffix == ".docx":
        return source
    if suffix != ".doc":
        raise RuntimeError(f"Unsupported template format: {source.suffix or '<none>'}")
    converted_path = Path(output_dir) / f"{source.stem}.docx"
    cmd = ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", str(Path(output_dir)), str(source)]
    _run_libreoffice(command=cmd, timeout=120, error_message="LibreOffice DOCX conversion timed out")
    if not converted_path.exists():
        raise RuntimeError(f"DOCX template not found after conversion: {converted_path}")
    return converted_path


def _resolve_recipient_label(has_kinship: bool, degree_of_kinship: str) -> str:
    if not has_kinship:
        return "себе"
    if degree_of_kinship == "ребенок":
        return "ребенку"
    if degree_of_kinship == "лицо, чьим законным представителем я являюсь":
        return "подопечному"
    return "родственнику"


def generate_begemotik_docx(
    template_path,
    iin,
    surname,
    name,
    last_name,
    gender,
    birthdate,
    phone,
    has_kinship,
    surname_kinship,
    name_kinship,
    last_name_kinship,
    degree_of_kinship,
    allergy_value,
    no_allergy_value,
    procedure,
    signature_base64,
    agreement_id,
    output_basename,
    output_dir,
) -> Path:
    normalized_template_path, suspicious_placeholders = _prepare_template_for_render(
        template_path=template_path, output_dir=output_dir, output_basename=output_basename,
    )
    tpl = DocxTemplate(normalized_template_path)
    sig_bytes = _decode_signature(signature_base64)
    sig_tmp = Path(output_dir) / f"{agreement_id}_sig.png"
    sig_tmp.write_bytes(sig_bytes)
    now = datetime.now()

    patient_full_name = " ".join(filter(None, [surname, name, last_name]))
    rep_full_name = " ".join(filter(None, [surname_kinship, name_kinship, last_name_kinship])) if has_kinship else ""
    name_surname_kinship = patient_full_name if has_kinship else ""
    if has_kinship:
        patient_field = ""
        kinship_field = rep_full_name
    else:
        patient_field = patient_full_name
        kinship_field = ""
    signature_kinship = rep_full_name if has_kinship else ""

    context = {
        "has_kinship": has_kinship,
        "representative_full_name": rep_full_name,
        "recipient_label": _resolve_recipient_label(has_kinship, degree_of_kinship),
        "iin": iin, "surname": surname, "name": name, "last_name": last_name,
        "gender": gender, "birthdate": birthdate, "phone": phone,
        "surname_kinship": surname_kinship, "name_kinship": name_kinship,
        "last_name_kinship": last_name_kinship, "degree_of_kinship": degree_of_kinship,
        "name_surname_kinship": name_surname_kinship, "signature_kinship": signature_kinship,
        "patient": patient_field, "kinship": kinship_field,
        "allergy": allergy_value, "no_allergy": no_allergy_value,
        "procedure": procedure,
        "date": now.strftime("%d.%m.%Y"), "full_date": now.strftime("%d.%m.%Y %H:%M"),
        "agreement_id": agreement_id,
        "signature": InlineImage(tpl, str(sig_tmp), width=Mm(50)),
    }
    try:
        tpl.render(context)
    except (TemplateSyntaxError, UndefinedError) as exc:
        raise RuntimeError(f"Template render error in {Path(template_path).name}: {exc}") from exc
    docx_path = Path(output_dir) / f"{output_basename}.docx"
    tpl.save(str(docx_path))
    return docx_path


def convert_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)]
    _run_libreoffice(command=cmd, timeout=120, error_message="LibreOffice conversion timed out")
    pdf_path = output_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError(f"PDF not found after conversion: {pdf_path}")
    return pdf_path
